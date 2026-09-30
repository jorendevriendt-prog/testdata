import csv, json, math, os, re, datetime as dt, zipfile
from collections import Counter
from pathlib import Path
from email import policy
from email.parser import BytesParser
import docx, openpyxl
from pypdf import PdfReader
from registry import DOCS, CONFIRM_CLUSTERS, CONTRA

ROOT = Path(__file__).resolve().parent.parent; D = ROOT / "documenten"
W_B = {"auteur": .20, "bron": .15, "datum": .10, "status": .20, "gebruik": .05, "tegenstellingen": .15, "bevestigingen": .15}
W_R = {"kernwoorden": .70, "actualiteit": .20, "gebruik": .10}

def extract(p: Path) -> str:
    s = p.suffix.lower()
    if s == ".docx":
        d = docx.Document(p)
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows: parts += [c.text for c in row.cells]
        return "\n".join(parts)
    if s == ".pdf":
        return "\n".join(pg.extract_text() or "" for pg in PdfReader(p).pages)
    if s == ".xlsx":
        wb = openpyxl.load_workbook(p, read_only=True)
        return "\n".join(" ".join(str(c) for c in row if c is not None) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    if s == ".eml":
        m = BytesParser(policy=policy.default).parse(open(p, "rb"))
        return m["Subject"] + "\n" + m.get_body(preferencelist=("plain",)).get_content()
    return p.read_text(encoding="utf-8")

STOP = set("""de het een en van in op te dat die dit voor met is zijn aan als bij of om door uit ook niet naar wordt worden werd kan kunnen moet
wij we je jij jouw uw u ze zij hij haar hem hun ons onze mijn ik er dan maar nog wel al tot na per over onder tussen zoals deze dus heeft hebben
geen meer alle elke iedereen wat wie waar hoe waarom welke zich zelf binnen tijdens vanaf vóór voor sinds zonder eerste twee drie volgens gelieve
the and via ter toe nu zo zal zullen was waren had heel even iets veel weinig goed goede nieuwe nieuw bijv eur euro dagen dag week jaar maand
arvelia services voskenslaan gent status versie pagina datum document opmerking waarde veld totaal aantal hr-be arvelia-hr
januari februari maart april juni juli augustus september oktober november december maandag dinsdag woensdag donderdag vrijdag
zaterdag zondag gaan weet krijgt zegt stuur kort beste correcte tijdige eerst minstens toepassing overeenkomstig betreffende
jullie iemand natuurlijk vertellen berichtje delen info update start intern taak toelichting inleiding afgerond gepland
omschrijving opgesteld behaald verslag correct tijd nooit ontdek kris mertens lotte peeters marc damme thomas pieter
sofie eline maes kevin ervaring ontvangen bedrag kantoor medewerker medewerkers team afdeling
hallo allemaal want anders mailbox bekijk browser vorige volgend klopt persoon maanden dienst blijft hierbij harte""".split())

def keywords(text, n=8):
    words = re.findall(r"[a-zà-ÿ][a-zà-ÿ0-9\-]{3,}", text.lower())
    c = Counter(w.strip("-") for w in words if w.strip("-") not in STOP and not w.isdigit())
    return [w for w, _ in c.most_common(n)]

def usage_score(n): return min(100, round(100 * math.log(1 + n) / math.log(1 + 400)))
def cat(x): return "hoog" if x >= 70 else ("midden" if x >= 40 else "laag")

def main():
    fname = {k: v[0] for k, v in DOCS.items()}
    conf = {k: set() for k in DOCS}; contra = {k: set() for k in DOCS}
    for members in CONFIRM_CLUSTERS.values():
        for a in members:
            for b in members:
                if a != b: conf[a].add(b)
    for a, b, _ in CONTRA: contra[a].add(b); contra[b].add(a)
    assert not any(conf[k] & contra[k] for k in DOCS), "paar zowel bevestigend als tegensprekend"

    meta, labels = [], []
    for k, v in DOCS.items():
        (f, dtype, auteur, functie, aanm, gew, status, versie, opens, bronnen, A, B, Dd, S, T, C, K, Ac, toel) = v
        p = D / f
        text = extract(p)
        g = usage_score(opens)
        sub_b = {"auteur": A, "bron": B, "datum": Dd, "status": S, "gebruik": g, "tegenstellingen": T, "bevestigingen": C}
        sub_r = {"kernwoorden": K, "actualiteit": Ac, "gebruik": g}
        bet = round(sum(W_B[x] * sub_b[x] for x in W_B)); rel = round(sum(W_R[x] * sub_r[x] for x in W_R))
        meta.append({"bestandsnaam": f, "auteur": auteur, "functie_auteur": functie, "aanmaakdatum": aanm, "laatst_gewijzigd": gew,
                     "status": status, "versie": versie, "aantal_keer_geopend": opens, "aantal_bronvermeldingen": bronnen,
                     "kernwoorden": keywords(text), "bevestigd_door": sorted(fname[x] for x in conf[k]),
                     "tegengesproken_door": sorted(fname[x] for x in contra[k])})
        labels.append({"bestandsnaam": f, "documenttype": dtype, "formaat": p.suffix.lstrip(".").lower(),
                       "betrouwbaarheid_score": bet, "relevantie_score": rel,
                       "betrouwbaarheid_categorie": cat(bet), "relevantie_categorie": cat(rel),
                       **{f"b_{x}": sub_b[x] for x in W_B}, **{f"r_{x}": sub_r[x] for x in W_R},
                       "aantal_bevestigingen": len(conf[k]), "aantal_tegenstellingen": len(contra[k]),
                       "toelichting": toel})
        # bestandsdatum = laatst gewijzigd
        y, m, d = map(int, gew.split("-")); ts = dt.datetime(y, m, d, 12).timestamp(); os.utime(p, (ts, ts))

    labels.sort(key=lambda r: r["bestandsnaam"].lower()); meta.sort(key=lambda r: r["bestandsnaam"].lower())
    for name, rows in (("labels", labels), ("metadata", meta)):
        (ROOT / f"{name}.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
        with open(ROOT / f"{name}.csv", "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), delimiter=";")
            w.writeheader()
            for r in rows: w.writerow({kk: (" | ".join(vv) if isinstance(vv, list) else vv) for kk, vv in r.items()})
    rel_rows = [{"document_a": fname[a], "document_b": fname[b], "type": "tegenstelling", "onderwerp": o} for a, b, o in CONTRA]
    for topic, members in CONFIRM_CLUSTERS.items():
        for i, a in enumerate(members):
            for b in members[i+1:]:
                rel_rows.append({"document_a": fname[a], "document_b": fname[b], "type": "bevestiging", "onderwerp": topic})
    with open(ROOT / "relaties.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["document_a", "document_b", "type", "onderwerp"], delimiter=";"); w.writeheader(); w.writerows(rel_rows)
    for r in labels:
        print(f"{r['bestandsnaam'][:52]:52} {r['formaat']:5} B={r['betrouwbaarheid_score']:3} R={r['relevantie_score']:3}")
    print("relaties:", len(rel_rows))

if __name__ == "__main__":
    main()
