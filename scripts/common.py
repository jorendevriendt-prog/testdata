# Gemeenschappelijke helpers en fictieve bedrijfsgegevens
import datetime as dt
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem

OUT = Path(__file__).resolve().parent.parent / "documenten"
OUT.mkdir(parents=True, exist_ok=True)

BEDRIJF = "Arvelia HR Services NV"
ADRES = "Voskenslaan 214, 9000 Gent"
KBO = "BE 0712.345.678"
IBAN = "BE71 0961 2345 6769"
BIC = "GKCCBEBB"
DOMEIN = "arvelia-hr.be"
TEL = "+32 9 123 45 67"
PC = "Paritair Comité 200 (aanvullend paritair comité voor bedienden)"

def dtm(y, m, d, h=10, mi=0):
    return dt.datetime(y, m, d, h, mi)

# ---------------- DOCX ----------------
def make_docx(filename, blocks, author="", title="", created=None, modified=None, header=None, footer=None, comments=""):
    from docx.enum.text import WD_COLOR_INDEX
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    if header:
        hp = sec.header.paragraphs[0]; hp.text = header
        hp.runs[0].font.size = Pt(8); hp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if footer:
        fp = sec.footer.paragraphs[0]; fp.text = footer
        fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    for b in blocks:
        k = b[0]
        if k == "title": d.add_heading(b[1], level=0)
        elif k == "h1": d.add_heading(b[1], level=1)
        elif k == "h2": d.add_heading(b[1], level=2)
        elif k == "p": d.add_paragraph(b[1])
        elif k == "pb":
            r = d.add_paragraph().add_run(b[1]); r.bold = True
        elif k == "pi":
            r = d.add_paragraph().add_run(b[1]); r.italic = True; r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif k == "red":
            r = d.add_paragraph().add_run(b[1]); r.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
        elif k == "hl":
            r = d.add_paragraph().add_run(b[1]); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif k == "bullets":
            for it in b[1]: d.add_paragraph(it, style="List Bullet")
        elif k == "num":
            for it in b[1]: d.add_paragraph(it, style="List Number")
        elif k == "table":
            rows = b[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for i, row in enumerate(rows):
                for j, val in enumerate(row): t.cell(i, j).text = str(val)
            d.add_paragraph()
        elif k == "sig":
            d.add_paragraph()
            t = d.add_table(rows=1, cols=len(b[1]))
            for j, s in enumerate(b[1]): t.cell(0, j).text = s
        elif k == "break": d.add_page_break()
    cp = d.core_properties
    cp.author = author; cp.last_modified_by = author; cp.title = title; cp.comments = comments
    created = created or dtm(2026, 1, 1)
    cp.created = created; cp.modified = modified or created; cp.revision = 1
    d.save(OUT / filename)

# ---------------- PDF ----------------
_ss = getSampleStyleSheet()
S = {
    "title": ParagraphStyle("t", parent=_ss["Title"], fontSize=17, spaceAfter=10, textColor=colors.HexColor("#1F3A5F")),
    "h1": ParagraphStyle("h1", parent=_ss["Heading2"], fontSize=12.5, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#1F3A5F")),
    "h2": ParagraphStyle("h2", parent=_ss["Heading3"], fontSize=11, spaceBefore=6, spaceAfter=3),
    "p": ParagraphStyle("p", parent=_ss["BodyText"], fontSize=9.8, leading=13.2, spaceAfter=5),
    "small": ParagraphStyle("s", parent=_ss["BodyText"], fontSize=8, leading=10, textColor=colors.grey),
    "pb": ParagraphStyle("pb", parent=_ss["BodyText"], fontSize=9.8, leading=13.2, fontName="Helvetica-Bold"),
    "red": ParagraphStyle("red", parent=_ss["BodyText"], fontSize=10, textColor=colors.HexColor("#B00000"), fontName="Helvetica-Bold"),
}

def make_pdf(filename, blocks, author="", title="", subject="", header=None, footer=None, watermark=None,
             created=None, modified=None, keywords=""):
    path = OUT / filename
    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5); canvas.setFillColor(colors.grey)
        if header: canvas.drawString(2*cm, A4[1]-1.2*cm, header)
        canvas.drawString(2*cm, 1.0*cm, (footer + "  |  " if footer else "") + f"Pagina {doc.page}")
        if watermark:
            canvas.setFont("Helvetica-Bold", 60); canvas.setFillColor(colors.Color(0.85, 0.2, 0.2, alpha=0.18))
            canvas.translate(A4[0]/2, A4[1]/2); canvas.rotate(40); canvas.drawCentredString(0, 0, watermark)
        canvas.restoreState()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=1.8*cm,
                            title=title, author=author, subject=subject, keywords=keywords, creator="")
    story = []
    for b in blocks:
        k = b[0]
        if k in S: story.append(Paragraph(b[1], S[k]))
        elif k == "bullets":
            story.append(ListFlowable([ListItem(Paragraph(x, S["p"]), leftIndent=12) for x in b[1]], bulletType="bullet", start="•", leftIndent=12))
            story.append(Spacer(1, 4))
        elif k == "table":
            rows = [[Paragraph(str(c), S["p"]) for c in r] for r in b[1]]
            t = Table(rows, colWidths=b[2] if len(b) > 2 else None, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                                   ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t); story.append(Spacer(1, 8))
        elif k == "space": story.append(Spacer(1, b[1]))
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    # PDF-datums nazetten met pypdf
    if created:
        from pypdf import PdfReader, PdfWriter
        r = PdfReader(str(path)); w = PdfWriter(clone_from=r)
        f = lambda x: x.strftime("D:%Y%m%d%H%M%S+02'00'")
        md = {k: v for k, v in (r.metadata or {}).items()}
        md.update({"/CreationDate": f(created), "/ModDate": f(modified or created), "/Producer": "Arvelia DMS PDF-export"})
        w.add_metadata(md)
        with open(path, "wb") as fh: w.write(fh)
