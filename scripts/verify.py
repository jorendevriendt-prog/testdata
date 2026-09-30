import csv, json, sys
from pathlib import Path
from build import extract, W_B, W_R, usage_score
import docx, openpyxl
from pypdf import PdfReader
ROOT = Path(__file__).resolve().parent.parent; D = ROOT / "documenten"
files = sorted(p.name for p in D.iterdir())
ok = True
for f in files:
    p = D / f
    t = extract(p)
    extra = ""
    if p.suffix == ".docx": cp = docx.Document(p).core_properties; extra = f"author={cp.author!r} created={cp.created:%Y-%m-%d}"
    if p.suffix == ".pdf": r = PdfReader(p); extra = f"pages={len(r.pages)} author={r.metadata.author!r} created={r.metadata.get('/CreationDate')}"
    if p.suffix == ".xlsx": wb = openpyxl.load_workbook(p); extra = f"sheets={wb.sheetnames} creator={wb.properties.creator!r}"
    if p.suffix == ".csv": extra = f"rows={len(list(csv.reader(open(p, encoding='utf-8'), delimiter=';')))}"
    words = len(t.split())
    if words < 80 and p.suffix != ".csv": ok = False
    print(f"{f[:50]:50} words={words:5} {extra}")
for name in ("labels", "metadata"):
    rows = list(csv.DictReader(open(ROOT / f"{name}.csv", encoding="utf-8-sig"), delimiter=";"))
    js = json.load(open(ROOT / f"{name}.json", encoding="utf-8"))
    names = [r["bestandsnaam"] for r in rows]
    cover = sorted(names) == files and len(set(names)) == len(names) == len(js) and [r["bestandsnaam"] for r in js] == names
    print(name, "rows", len(rows), "exact cover:", cover); ok &= cover
lab = json.load(open(ROOT / "labels.json")); meta = {m["bestandsnaam"]: m for m in json.load(open(ROOT / "metadata.json"))}
for r in lab:
    b = round(sum(W_B[k] * r[f"b_{k}"] for k in W_B)); rl = round(sum(W_R[k] * r[f"r_{k}"] for k in W_R))
    m = meta[r["bestandsnaam"]]
    assert b == r["betrouwbaarheid_score"] and rl == r["relevantie_score"]
    assert r["b_gebruik"] == usage_score(m["aantal_keer_geopend"])
    for o in m["bevestigd_door"]: assert r["bestandsnaam"] in meta[o]["bevestigd_door"]
    for o in m["tegengesproken_door"]: assert r["bestandsnaam"] in meta[o]["tegengesproken_door"]
    for o in m["bevestigd_door"] + m["tegengesproken_door"]: assert o in files
print("scores reproduceerbaar, relaties symmetrisch, verwijzingen geldig")
print("ALLES OK" if ok else "PROBLEEM"); sys.exit(0 if ok else 1)
